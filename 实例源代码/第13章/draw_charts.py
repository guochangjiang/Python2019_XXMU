# draw_charts.py
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.text import CharacterProperties

# 读取数据
rows = []
with open("GDP.tsv", 'r', encoding='UTF-8') as DATA:
    for line in DATA:
        line = line.strip()
        if line:
            rows.append(line.split("\t"))
# 将数据转换为数值类型
for i in range(1, len(rows)):
    for j in range(1, len(rows[i])):
        rows[i][j] = eval(rows[i][j])

wb = Workbook(write_only=True)
ws = wb.create_sheet()
for row in rows:
    ws.append(row)
chart1 = BarChart()
chart1.type = "col"
chart1.style = 10
chart1.title = "柱状图"
chart1.y_axis.title = 'GDP(万亿美元)'
chart1.x_axis.title = '年份'
# 设置标题字号
cp = CharacterProperties(sz=1200)
chart1.title.tx.rich.p[0].r.rPr = cp
chart1.y_axis.title.tx.rich.p[0].r.rPr = cp
chart1.x_axis.title.tx.rich.p[0].r.rPr = cp

data = Reference(ws, min_col=2, min_row=1, max_row=6, max_col=4)
cats = Reference(ws, min_col=1, min_row=2, max_row=6)
chart1.add_data(data, titles_from_data=True)
chart1.set_categories(cats)
chart1.shape = 4
chart1.width = 12
chart1.height = 6
ws.add_chart(chart1, "A8")

wb.save("bar.xlsx")